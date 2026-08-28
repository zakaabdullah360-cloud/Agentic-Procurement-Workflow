"""
Tool 2: Purchase Order Generator. Produces a real .docx file on disk (not a
mockup) plus the structured PO dict used for validation/approval.
"""
import datetime as dt
import os

from docx import Document

from app.config import GENERATED_PO_DIR


def build_po_dict(workflow_id: int, structured: dict, best_quote: dict) -> dict:
    return {
        "workflow_id": workflow_id,
        "item": structured["item"],
        "quantity": best_quote["quantity"],
        "unit_price": best_quote["unit_price"],
        "subtotal": round(best_quote["unit_price"] * best_quote["quantity"], 2),
        "total": round(best_quote["unit_price"] * best_quote["quantity"], 2),
        "currency": best_quote["currency"],
        "supplier_name": best_quote["supplier_name"],
        "supplier_id": best_quote["supplier_id"],
        "terms": "Net 30. Delivery/activation within "
                 f"{best_quote['lead_time_days']} days of PO approval. "
                 f"Warranty/contract term: {best_quote['warranty_years']} year(s).",
        "date": dt.date.today().isoformat(),
    }


def generate_po_document(po: dict) -> str:
    """Writes a .docx file and returns its path."""
    doc = Document()
    doc.add_heading("Purchase Order", level=1)
    doc.add_paragraph(f"PO Reference: WF-{po['workflow_id']}-{dt.date.today().strftime('%Y%m%d')}")
    doc.add_paragraph(f"Date: {po['date']}")
    doc.add_paragraph(f"Supplier: {po['supplier_name']}")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Item", "Quantity", "Unit Price", "Subtotal"
    row = table.add_row().cells
    row[0].text = po["item"]
    row[1].text = str(po["quantity"])
    row[2].text = f"{po['unit_price']:,.2f} {po['currency']}"
    row[3].text = f"{po['subtotal']:,.2f} {po['currency']}"

    doc.add_paragraph("")
    doc.add_paragraph(f"Total: {po['total']:,.2f} {po['currency']}").bold = True
    doc.add_paragraph(f"Terms: {po['terms']}")
    doc.add_paragraph("Status: Pending Approval")

    filename = f"PO_workflow_{po['workflow_id']}.docx"
    path = os.path.join(GENERATED_PO_DIR, filename)
    doc.save(path)
    return path
